import os
from typing import List
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from core.models import ResultadoAuditoria


VERDE = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
ROJO = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
GRIS = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
NEGRITA = Font(bold=True)
BORDE = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

COLUMNAS = [
    "Fecha Procesamiento", "Archivo", "Proveedor", "Región", "Tipo", "Resultado",
    "Fecha Emisión", "Fecha Vencimiento", "Visita 1", "Visita 2",
    "Distribuidor", "Característica 1", "Característica 2", "Característica 3",
    "Tipo Entrega", "DNI", "Nombre", "Apellido", "Tipo Vínculo", "Firma",
    "Errores",
]

ANCHOS = [16, 30, 12, 10, 14, 12, 14, 16, 14, 14, 25, 20, 20, 20, 14, 14, 16, 16, 14, 8, 50]


def _formato_fecha(d):
    return d.strftime("%d/%m/%Y") if d else ""


def _crear_encabezados(ws):
    for col, nombre in enumerate(COLUMNAS, start=1):
        celda = ws.cell(row=1, column=col, value=nombre)
        celda.font = NEGRITA
        celda.fill = GRIS
        celda.alignment = Alignment(horizontal="center")
        celda.border = BORDE
    for col, ancho in enumerate(ANCHOS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = ancho


def actualizar_excel(resultados: List[ResultadoAuditoria], carpeta_salida: str, fecha_procesamiento: str) -> str:
    os.makedirs(carpeta_salida, exist_ok=True)
    ruta = os.path.join(carpeta_salida, "reporte.xlsx")

    if os.path.exists(ruta):
        wb = load_workbook(ruta)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Auditoría"
        _crear_encabezados(ws)

    fila_inicio = ws.max_row + 1

    for fila, r in enumerate(resultados, start=fila_inicio):
        e = r.extraccion
        visita1 = _formato_fecha(e.visitas[0].fecha) if len(e.visitas) > 0 else ""
        visita2 = _formato_fecha(e.visitas[1].fecha) if len(e.visitas) > 1 else ""
        carac = e.caracteristicas_casa + [""] * (3 - len(e.caracteristicas_casa))
        errores_texto = "; ".join(f"{err.campo}: {err.mensaje}" for err in r.errores)

        valores = [
            fecha_procesamiento,
            r.archivo,
            r.proveedor,
            r.region,
            r.tipo_documento or "",
            "APROBADO" if r.aprobado else "RECHAZADO",
            _formato_fecha(e.fecha_emision),
            _formato_fecha(e.fecha_vencimiento),
            visita1,
            visita2,
            e.distribuidor or "",
            carac[0], carac[1], carac[2],
            e.tipo_entrega or "",
            e.dni or "",
            e.nombre or "",
            e.apellido or "",
            e.tipo_vinculo or "",
            "Sí" if e.tiene_firma else "No",
            errores_texto,
        ]

        fill = VERDE if r.aprobado else ROJO
        for col, valor in enumerate(valores, start=1):
            celda = ws.cell(row=fila, column=col, value=valor)
            celda.fill = fill
            celda.border = BORDE
            celda.alignment = Alignment(wrap_text=True)

    wb.save(ruta)
    return ruta
